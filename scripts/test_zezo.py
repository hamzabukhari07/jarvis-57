"""
scripts/test_zezo.py — Interactive Test Runner & Feature Inspector for ZEZO OS

Use this script to test all newly integrated features:
1. Universal Multi-Format File Ingestion (Word, PDF, Scanned PDF, Vision, Truncation, Legacy Rejection)
2. Live Mobile Remote Web Dashboard & Telemetry
3. Launch Full Voice Assistant (main.py)

Architect: Hamza Bukhari
"""

import os
import sys
import time
import json
import socket
from pathlib import Path

# UTF-8 encoding safeguard for Windows consoles
for _stream in ("stdout", "stderr"):
    try:
        _s = getattr(sys, _stream, None)
        if _s is not None and hasattr(_s, "reconfigure"):
            _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

# Add project root to sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from core.file_reader import read_file
from actions.file_processor import file_processor


def _ensure_test_files():
    """Generate all 6 real runtime test files in D:/test."""
    os.makedirs("D:/test", exist_ok=True)
    from PIL import Image, ImageDraw

    # 1. DOCX
    try:
        from docx import Document
        doc = Document()
        doc.add_heading("ZEZO Test Document", 0)
        doc.add_paragraph("This is a verified test document for autonomous file ingestion.")
        doc.save("D:/test/sample.docx")
    except Exception as e:
        print("[!] Note on DOCX creation:", e)

    # 2. Text PDF (clean native text via reportlab)
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas
        c = canvas.Canvas("D:/test/report.pdf", pagesize=letter)
        c.drawString(100, 750, "ZEZO Autonomous OS Annual Architecture Report")
        c.drawString(100, 720, "Executive Summary: The system operates at sub-500ms latency across voice and automation.")
        c.drawString(100, 690, "Section 1: The unified file reader processes Word, PDF, Excel, and code files.")
        c.save()
    except Exception as e:
        print("[!] Note on PDF creation:", e)

    # 3. Scanned PDF (Pure image raster with no text stream)
    try:
        img = Image.new("RGB", (800, 600), color=(255, 255, 255))
        draw = ImageDraw.Draw(img)
        draw.text((50, 50), "INVOICE #98231\nClient: Hamza Bukhari\nAmount: $3,250.00\nStatus: Paid in Full", fill=(0, 0, 0))
        img_path = "D:/test/temp_scan.png"
        img.save(img_path)
        img_pdf = Image.open(img_path)
        img_pdf.save("D:/test/scanned.pdf", "PDF", resolution=100.0)
        if os.path.exists(img_path):
            os.remove(img_path)
    except Exception as e:
        print("[!] Note on Scanned PDF creation:", e)

    # 4. Screenshot PNG
    try:
        screenshot = Image.new("RGB", (600, 400), color=(15, 15, 20))
        sdraw = ImageDraw.Draw(screenshot)
        sdraw.rectangle([20, 20, 580, 80], fill=(242, 78, 30))
        sdraw.text((40, 40), "ZEZO HUD TELEMETRY · CPU: 12% | RAM: 34%", fill=(255, 255, 255))
        sdraw.text((40, 120), "Autonomous Desktop OS v2\nTerminal Status: Operational", fill=(200, 200, 200))
        screenshot.save("D:/test/screenshot.png")
    except Exception as e:
        print("[!] Note on Image creation:", e)

    # 5. Large JSON (3MB for 64KB truncation guard)
    data = {"items": ["x" * 1000 for _ in range(3000)]}
    with open("D:/test/large.json", "w") as f:
        json.dump(data, f)

    # 6. Legacy old.doc
    with open("D:/test/old.doc", "wb") as f:
        f.write(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1 Legacy Word Binary Compound File")


def test_file_reader_suite():
    """Run interactive tests on all 6 file types."""
    _ensure_test_files()
    print("\n" + "═" * 65)
    print(" 🚀 ZEZO MULTI-FORMAT FILE INGESTION — RUNTIME TESTS")
    print("═" * 65)

    tests = [
        ("1. Modern Word DOCX", "D:/test/sample.docx", "summarize"),
        ("2. Text-Based PDF", "D:/test/report.pdf", "summarize"),
        ("3. Scanned Image PDF (OCR)", "D:/test/scanned.pdf", "summarize"),
        ("4. UI Screenshot (Vision)", "D:/test/screenshot.png", "describe"),
        ("5. 3MB JSON (64KB Guard)", "D:/test/large.json", "read"),
        ("6. Legacy .DOC (Rejection)", "D:/test/old.doc", "summarize"),
    ]

    for name, path, action in tests:
        print(f"\n▶ Running: {name}")
        print(f"  Target: {path} | Action: {action}")
        t0 = time.perf_counter()
        
        if action == "read":
            res = read_file(path)
            output = res.text
            engine = res.engine
            trunc = res.is_truncated
        else:
            output = file_processor({"file_path": path, "action": action})
            engine = "file_processor"
            trunc = len(output) > 65536

        t1 = time.perf_counter()
        print(f"  ⏱ Time: {t1 - t0:.2f}s | Engine: {engine} | Truncated: {trunc}")
        print("  ┌─ Output Preview ──────────────────────────────────────────┐")
        for line in str(output)[:350].split("\n"):
            print(f"  │ {line}")
        if len(str(output)) > 350:
            print("  │ ... [truncated in preview]")
        print("  └───────────────────────────────────────────────────────────┘")


def get_local_ip():
    """Get local network IP for mobile web dashboard access."""
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
    except Exception:
        ip = "127.0.0.1"
    finally:
        s.close()
    return ip


def launch_dashboard():
    """Start standalone dashboard server."""
    from dashboard import server
    ip = get_local_ip()
    port = 8765
    print("\n" + "═" * 65)
    print(" 📱 ZEZO MOBILE REMOTE ACCESS DASHBOARD")
    print("═" * 65)
    print(f" Local URL:   http://localhost:{port}")
    print(f" Mobile URL:  http://{ip}:{port}")
    print(" Theme:       Autonomous Traffic Vectors (Matte Void #050505 + Cyber Orange #f24e1e)")
    print(" Features:    Voice Streaming (PCM16), Live Telemetry, File Drag-and-Drop")
    print(" Press Ctrl+C to stop the dashboard.")
    print("═" * 65 + "\n")
    server.start_server(host="0.0.0.0", port=port, blocking=True)


def main():
    print("""
  ███████╗███████╗███████╗ ██████╗ 
  ╚══███╔╝██╔════╝╚══███╔╝██╔═══██╗
    ███╔╝ █████╗    ███╔╝ ██║   ██║
   ███╔╝  ██╔══╝   ███╔╝  ██║   ██║
  ███████╗███████╗███████╗╚██████╔╝
  ╚══════╝╚══════╝╚══════╝ ╚═════╝ 
  Autonomous Desktop AI Operating System v2
  Lead Architect: Hamza Bukhari
    """)
    print("Select an option:")
    print(" 1. Run Universal File Reader Runtime Tests (Word, PDF, Scanned, Vision, 64KB Guard)")
    print(" 2. Start Mobile Remote Access Web Dashboard (Autonomous Traffic Vectors UI)")
    print(" 3. Launch Full ZEZO Voice Operating System (main.py)")
    print(" 4. Run Automated Pytest Suites (verify 100% pass)")
    print(" 0. Exit")

    choice = input("\nEnter choice [1-4, 0]: ").strip()

    if choice == "1":
        test_file_reader_suite()
    elif choice == "2":
        launch_dashboard()
    elif choice == "3":
        print("\n🚀 Launching ZEZO OS via main.py...")
        import main as zezo_main
        zezo_main.main()
    elif choice == "4":
        import subprocess
        print("\n🧪 Running Pytest Verification Suite...")
        subprocess.run(["python", "-m", "pytest", "tests/test_file_reader_suite.py", "tests/test_remote_dashboard_suite.py", "-v"])
    else:
        print("Exiting.")


if __name__ == "__main__":
    main()
